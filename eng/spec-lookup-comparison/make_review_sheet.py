# SPDX-License-Identifier: Apache-2.0
# Licensed to the Ed-Fi Alliance under one or more agreements.
# The Ed-Fi Alliance licenses this file to you under the Apache License, Version 2.0.
# See the LICENSE and NOTICES files in the project root for more information.

"""Blind second-review sheet for the comparison runs: every answer whose score was adjusted or is below 2,
plus a random sample of answers scored 2, shuffled, without the condition.

Usage: python make_review_sheet.py <results-dir> [--sample 13] [--seed 148]
Writes <results-dir>/second-review.xlsx (to share) and <results-dir>/second-review-key.json (keep; maps each ID back
to condition, run and first-pass score).
"""
import argparse
import json
import os
import random

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

RUBRIC = [
    ('2', 'Correct and complete: names the expected resource(s) or fields. Extra correct detail is fine.'),
    ('1', 'Partially correct: right resource but incomplete or partly wrong key/fields, or the right answer '
          'buried among wrong alternatives.'),
    ('0', 'Wrong, missing, or only "I don\'t know".'),
]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('results')
    ap.add_argument('--sample', type=int, default=13)
    ap.add_argument('--seed', type=int, default=148)
    args = ap.parse_args()
    scores = json.load(open(os.path.join(args.results, 'scores_final.json'), encoding='utf-8'))
    runs = {(r['condition'], r['task'], r['run']): r
            for r in map(json.loads, open(os.path.join(args.results, 'runs.jsonl'), encoding='utf-8'))}
    rng = random.Random(args.seed)
    flagged = [s for s in scores if s['adjusted'] or s['correctness_final'] < 2
               or s['correctness'] != s['correctness_final']]
    rest = [s for s in scores if s not in flagged]
    picked = flagged + rng.sample(rest, args.sample)
    rng.shuffle(picked)

    wb = Workbook()
    info = wb.active
    info.title = 'Instructions'
    lines = [
        'Ed-Fi spec lookup comparison : second review',
        '',
        'Each row is one assistant answer to an Ed-Fi API (Data Standard 5.2) lookup question. Which setup '
        'produced it is hidden.',
        'Score each answer in the "Score" column of the Answers sheet, comparing it with the expected answer '
        '(verified against the specification).',
        'Some answers mention reading files or an index; please score the content only.',
        '',
        'Rubric:',
    ] + [f'  {k} = {v}' for k, v in RUBRIC] + [
        '',
        'Notes are optional; add one when the score is below 2 or you are unsure.',
    ]
    for i, line in enumerate(lines, 1):
        info.cell(i, 1, line).font = Font(bold=(i == 1))
    info.column_dimensions['A'].width = 120

    ws = wb.create_sheet('Answers')
    head = ['ID', 'Question', 'Expected answer', 'Answer', 'Score (0-2)', 'Notes']
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill('solid', fgColor='DDEBF7')
    key = {}
    for n, s in enumerate(picked, 1):
        rid = f'R{n:02d}'
        r = runs[(s['condition'], s['task'], s['run'])]
        ws.append([rid, r['prompt'], r['expected'], r['answer'], None, None])
        key[rid] = {k: s[k] for k in ('condition', 'task', 'run', 'correctness', 'correctness_final')} | {
            'flagged': s in flagged}
    for col, width in zip('ABCDEF', (7, 40, 40, 100, 12, 40)):
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical='top')
    dv = DataValidation(type='list', formula1='"0,1,2"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f'E2:E{len(picked) + 1}')
    ws.freeze_panes = 'B2'

    wb.save(os.path.join(args.results, 'second-review.xlsx'))
    with open(os.path.join(args.results, 'second-review-key.json'), 'w', encoding='utf-8') as fh:
        json.dump(key, fh, indent=1)
    print(f'{len(picked)} answers ({len(flagged)} flagged, {args.sample} sampled)')


if __name__ == '__main__':
    main()
